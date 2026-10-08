"""Regression tests for the bugs confirmed in the project-wide code review."""

import asyncio
import contextlib
import html
import types

import openpyxl
import pandas as pd
import pytest
from git.exc import GitCommandError

from functions.base.base import BaseFunction


def ns(**kw):
    return types.SimpleNamespace(**kw)


def collect():
    msgs = []

    async def report(text):
        msgs.append(text)

    return msgs, report


class _Storage:
    initialize = False

    def __init__(self, sessions=()):
        self.sessions = list(sessions)
        self.forgotten = []
        self.jsessions_paths = {}

    @contextlib.asynccontextmanager
    async def ainitialize_session(self, session):
        yield

    def forget_session(self, path):
        self.forgotten.append(path)

    def get_session_path(self, session):
        return None

    def remember_username(self, session, username):
        pass

    def remember_name(self, session, first_name, last_name):
        pass


async def _no_sleep(*_):
    return None


# --- functions/ --------------------------------------------------------------

class TestJoinerCaptcha:
    def test_handler_only_no_run_until_disconnected(self):
        from functions.joiner import JoinerFunc

        class _S:
            handlers = []

            def add_event_handler(self, cb, ev):
                self.handlers.append(cb)

            async def run_until_disconnected(self):  # would disconnect on cancel
                raise AssertionError("must not be used")

        fn = JoinerFunc(_Storage(), ns(delay=[0]))
        s = _S()
        assert fn.solve_captcha(s) is None  # plain call, no task to cancel
        assert s.handlers == [fn.on_message]


class TestCtrlCCancelsFunction:
    def test_task_cancelled_and_sessions_reconnected(self):
        from modules.storages.functions_storage import FunctionsStorage

        loop = asyncio.new_event_loop()
        cancelled = []

        class _Client:
            connected = False

            def is_connected(self):
                return self.connected

            async def connect(self):
                self.connected = True

        client = _Client()

        async def listener():
            try:
                await asyncio.sleep(10)
            except asyncio.CancelledError:
                cancelled.append(True)
                raise

        def _interrupt():
            raise KeyboardInterrupt

        fs = FunctionsStorage.__new__(FunctionsStorage)
        fs.loop = loop
        fs.storage = ns(initialize=True, sessions=[client])
        fs.functions = [(ns(execute=listener), "doc")]

        loop.call_later(0.01, _interrupt)  # Ctrl-C while the loop idles, not inside the task
        try:
            fs.run_instance(fs.functions[0][0])
            assert False, "KeyboardInterrupt must propagate to the menu"
        except KeyboardInterrupt:
            pass
        finally:
            asyncio.set_event_loop(None)  # execute() set it; don't leak a closed loop
            loop.close()

        assert cancelled == [True]  # not left pending for the next menu function
        assert client.connected is True


class TestChangeNameBlankLines:
    def test_blank_lines_dropped(self, monkeypatch, tmp_path):
        from functions.changename import ChangeNameFunc

        (tmp_path / "assets").mkdir()
        (tmp_path / "assets" / "names.txt").write_text("Ivan\n\nPetr Petrov\n", encoding="utf-8")
        monkeypatch.chdir(tmp_path)
        monkeypatch.setattr("functions.changename.console.input", lambda *_: "y")

        fn = ChangeNameFunc(_Storage(), ns(delay=[0], profile_pause=[0]))
        fn.ask_accounts_count = lambda: None
        got = []

        async def run(report, names=None, **_):
            got.append(names)

        fn.run = run
        asyncio.run(fn.execute())

        assert got == [["Ivan", "Petr Petrov"]]


class TestSessionsStorage:
    def test_binary_session_file_is_skipped(self, tmp_path):
        from modules.storages.sessions_storage import SessionsStorage

        (tmp_path / "sqlite.session").write_bytes(b"SQLite format 3\x00\xff\xfe\x80")
        storage = SessionsStorage(str(tmp_path), 1, "hash", initialize=False)
        assert storage.sessions == []

    def test_authorization_check_error_skips_session(self, tmp_path):
        from modules.storages.sessions_storage import SessionsStorage

        storage = SessionsStorage(str(tmp_path), 1, "hash", initialize=False)

        class _S:
            async def connect(self):
                return None

            async def is_user_authorized(self):
                raise RuntimeError("network")

            async def disconnect(self):
                return None

        s = _S()
        storage.full_sessions["p"] = s
        asyncio.run(storage.check_session(s, "p"))  # must not raise into the init gather
        assert storage.sessions == []


class TestReportCli:
    def test_rest_goes_through_replay_rest(self, monkeypatch):
        from functions.report import ReportFunc

        class _S:
            async def get_me(self):
                return ns(first_name="A")

        first, other = _S(), _S()
        fn = ReportFunc(_Storage([first, other]), ns(delay=[0]))
        fn.ask_accounts_count = lambda: None
        answers = iter(["https://t.me/x", "1"])
        monkeypatch.setattr("functions.report.Prompt.ask", lambda *a, **k: next(answers))
        monkeypatch.setattr("functions.report.console.input", lambda *_: "c")

        async def resolve_and_report(*_):
            return [0]

        replayed = []

        async def replay_rest(sessions, *args):
            replayed.append(sessions)

        fn.resolve_and_report = resolve_and_report
        fn.replay_rest = replay_rest
        asyncio.run(fn.execute())

        assert replayed == [[other]]


class TestCommentsErrorReset:
    def test_success_resets_counter(self, monkeypatch):
        from functions.broadcast_comments import CommentsBroadcastFunc

        outcomes = iter([ValueError("x")] * 4 + [None] + [ValueError("x")] * 4 + [None])
        sends = []

        async def fake_send(*a, **k):
            sends.append(1)
            err = next(outcomes)
            if err is not None:
                raise err

        monkeypatch.setattr("functions.broadcast_comments.rich_message.send", fake_send)

        class _S:
            async def get_me(self):
                return ns(first_name="A")

        fn = CommentsBroadcastFunc(_Storage(), ns(delay=[0], messages_count=2))
        fn.delay = _no_sleep
        _, report = collect()
        asyncio.run(fn.broadcast(_S(), "chan", 1, None, report))

        assert len(sends) == 10  # 4 errors never reach the 5-error stop once reset


class TestSpamBlockMove:
    def test_moved_session_is_forgotten(self, monkeypatch, tmp_path):
        from functions.spamblock import SpamBlockFunc

        monkeypatch.chdir(tmp_path)
        (tmp_path / "sessions").mkdir()
        (tmp_path / "sessions" / "a.session").write_text("x")
        path = "sessions/a.session"

        storage = _Storage()
        storage.get_session_path = lambda s: path
        class _S:
            disconnected = False

            async def disconnect(self):
                self.disconnected = True

        session = _S()
        fn = SpamBlockFunc(storage, ns(delay=[0]))
        asyncio.run(fn.move_restricted({"permanent": [session]}))

        assert (tmp_path / "sessions" / "restricted" / "permanent" / "a.session").exists()
        assert storage.forgotten == [path]
        assert session.disconnected  # CLI clients stay connected: a forgotten one would run on

    def test_unblock_failure_is_reported(self):
        from telethon.errors import YouBlockedUserError
        from functions.spamblock import SpamBlockFunc

        class _S:
            async def get_me(self):
                return ns(username="C1", first_name="Acc", last_name=None)

            def conversation(self, *_):
                raise YouBlockedUserError(request=None)

            async def __call__(self, request):
                raise RuntimeError("no unblock")

        fn = SpamBlockFunc(_Storage(), ns(delay=[0]))
        msgs, report = collect()
        assert asyncio.run(fn.check(_S(), report)) is None
        assert any("⚠️ @C1 — не удалось разблокировать @SpamBot" in m for m in msgs)


class TestSpamBlockReport:
    def _check(self, reply, username="C1", progress=None):
        from functions.spamblock import SpamBlockFunc

        class _Conv:
            async def __aenter__(self):
                return self

            async def __aexit__(self, *exc):
                return False

            async def send_message(self, text):
                pass

            async def get_response(self):
                return ns(message=reply)

        class _S:
            async def get_me(self):
                return ns(id=7, username=username, first_name="Acc", last_name=None)

            def conversation(self, *_):
                return _Conv()

        msgs, report = collect()
        func = SpamBlockFunc(_Storage(), ns(delay=[0]))
        func.progress = progress
        result = asyncio.run(func.check(_S(), report))
        return msgs, result

    def test_restricted_workers_count_as_errors_in_the_summary(self):
        from bot.services.progress import Progress

        progress = Progress()
        self._check("Good news, no limits are currently applied.", progress=progress)
        self._check("Unfortunately...\nlimited until 12 Nov 2026, 10:00 UTC.", progress=progress)
        self._check("К сожалению...\nВаш аккаунт ограничен.", progress=progress)
        assert (progress.ok, progress.failed) == (1, 2)

    def test_active_names_the_account(self):
        msgs, result = self._check("Good news, no limits are currently applied.")
        assert msgs == ["✅ @C1 — без ограничений"]
        assert result[0] == "active"

    def test_restricted_until_names_the_account(self):
        msgs, result = self._check("Unfortunately...\nlimited until 12 Nov 2026, 10:00 UTC.")
        assert msgs == ["🚫 @C1 — ЛС ограничены до 12 Nov 2026"]
        assert result[0] == "12 Nov 2026"

    def test_falls_back_to_name_without_username(self):
        msgs, _ = self._check("Good news, no limits are currently applied.", username=None)
        assert msgs == ["✅ Acc — без ограничений"]

    def test_a_date_in_any_language_is_a_date(self):
        # @SpamBot answers in the session's language (random for an account added by phone)
        for reply, date in [
            ("К сожалению...\nОграничения будут сняты 12 нояб. 2026 г., 10:00 UTC.", "12 нояб. 2026"),
            ("很遗憾……\n您的账号将于 2026年11月12日 10:00 UTC 解除限制。", "2026年11月12日"),
            ("Niestety...\nOgraniczenia zostaną zniesione 12.11.2026, 10:00 UTC.", "12.11.2026"),
            ("Unfortunately...\nreleased in 2026, somehow.", "2026"),  # a year alone: still a date
        ]:
            msgs, result = self._check(reply)
            assert result[0] == date, reply
            assert msgs == [f"🚫 @C1 — ЛС ограничены до {date}"]

    def test_no_year_is_permanent(self):
        msgs, result = self._check("К сожалению...\nВаш аккаунт ограничен.")
        assert result[0] == "permanent"
        assert msgs == ["⛔ @C1 — ограничен бессрочно"]


class TestSpamBlockUnusableWorkers:
    """Dead sessions leave the pool; permanently restricted ones are excluded until clean."""

    class _S:
        def __init__(self, path, reply=None, me=True):
            self.path, self.reply, self.me = path, reply, me
            self.session = ns(_entities=set())

        async def connect(self):
            pass

        async def disconnect(self):
            pass

        async def get_me(self):
            if self.me is None:
                return None  # what Telethon returns for a banned / logged-out account
            if isinstance(self.me, Exception):
                raise self.me
            return ns(id=7, username="C1", first_name="Acc", last_name=None)

        def conversation(self, *_):
            reply = self.reply

            class _Conv:
                async def __aenter__(self):
                    return self

                async def __aexit__(self, *exc):
                    return False

                async def send_message(self, text):
                    pass

                async def get_response(self):
                    return ns(message=reply)

            return _Conv()

    def _storage(self, tmp_path, sessions):
        from modules.storages.sessions_storage import SessionsStorage

        (tmp_path / "sessions").mkdir(exist_ok=True)
        storage = SessionsStorage("sessions", 1, "x", initialize=False)
        for s in sessions:
            (tmp_path / s.path).write_text("x")
            storage.full_sessions[s.path] = s
        return storage

    def _run(self, storage, replies=False):
        from functions.spamblock import SpamBlockFunc

        fn = SpamBlockFunc(storage, ns(delay=[0]))
        fn.sessions = storage.sessions
        msgs, report = collect()
        asyncio.run(fn.run(report, replies=replies))
        return msgs

    def test_spambot_replies_are_shown_once_each(self, monkeypatch, tmp_path):
        monkeypatch.chdir(tmp_path)
        forever = "Unfortunately...\nyou are limited forever."
        workers = [self._S(f"sessions/{n}.jsession", reply=forever) for n in "ab"]
        workers.append(self._S("sessions/c.jsession", reply="Good news, no limits are currently applied."))
        storage = self._storage(tmp_path, workers)

        assert not any(m.startswith("💬") for m in self._run(storage))  # only when asked

        quotes = [m for m in self._run(storage, replies=True) if m.startswith("💬")]
        assert quotes == [f"💬 Ответ @SpamBot (@C1, @C1):\n{forever}"]  # the clean one: no quote

    def test_spambot_reply_is_cut_to_700(self, monkeypatch, tmp_path):
        monkeypatch.chdir(tmp_path)
        storage = self._storage(tmp_path, [self._S("sessions/a.jsession", reply="x\n" + "y" * 2000)])

        quote = [m for m in self._run(storage, replies=True) if m.startswith("💬")][0]
        assert quote.endswith("…") and len(quote.split(":\n", 1)[1]) == 701

    def test_entries_of_gone_sessions_are_dropped(self, monkeypatch, tmp_path):
        from modules import restricted_workers

        monkeypatch.chdir(tmp_path)
        restricted_workers.save(["sessions/gone.jsession"])
        restricted_workers.save_status({"sessions/gone2.jsession": "active"})
        s = self._S("sessions/a.jsession", reply="Good news, no limits are currently applied.")

        self._run(self._storage(tmp_path, [s]))

        assert restricted_workers.load() == []
        assert restricted_workers.load_status() == {"sessions/a.jsession": "active"}

    def test_dead_session_moves_to_inactive(self, monkeypatch, tmp_path):
        monkeypatch.chdir(tmp_path)
        dead = self._S("sessions/d.jsession", me=None)
        alive = self._S("sessions/a.jsession", reply="Good news, no limits are currently applied.")
        storage = self._storage(tmp_path, [dead, alive])

        msgs = self._run(storage)

        assert (tmp_path / "sessions" / "inactive" / "d.jsession").exists()
        assert storage.sessions == [alive]
        assert any(m.startswith("💀 d.jsession — сессия мертва") for m in msgs)

    def test_permanent_is_excluded_until_clean(self, monkeypatch, tmp_path):
        from modules import contacts_ledger, restricted_workers

        monkeypatch.chdir(tmp_path)
        contacts_ledger.save({"1": 7, "2": 7, "3": 8})
        s = self._S("sessions/a.jsession", reply="Unfortunately...\nyou are limited forever.")
        storage = self._storage(tmp_path, [s])

        msgs = self._run(storage)
        assert restricted_workers.load() == ["sessions/a.jsession"]
        assert storage.sessions == [s]  # stays in sessions/
        assert "⛔ @C1 — ограничен бессрочно, его контакты (2) переданы другим воркерам" in msgs  # no .jsession: shared

        s.me = RuntimeError("network")  # a failed check keeps the entry
        self._run(storage)
        assert restricted_workers.load() == ["sessions/a.jsession"]

        s.me, s.reply = True, "Good news, no limits are currently applied."
        self._run(storage)
        assert restricted_workers.load() == []

    def test_report_is_grouped_working_first(self, monkeypatch, tmp_path):
        monkeypatch.chdir(tmp_path)
        workers = [  # every group out of place
            self._S("sessions/e.jsession", me=RuntimeError("network")),
            self._S("sessions/d.jsession", me=None),
            self._S("sessions/c.jsession", reply="Unfortunately...\nyou are limited forever."),
            self._S("sessions/b.jsession", reply="Unfortunately...\nyou are limited until 12 Nov 2026."),
            self._S("sessions/a.jsession", reply="Good news, no limits are currently applied."),
        ]
        msgs = self._run(self._storage(tmp_path, workers))

        assert [m.split()[0] for m in msgs] == ["✅", "🚫", "⛔", "💀", "⚠️"]
        assert msgs[-1] == "⚠️ e.jsession — не удалось опросить: network"
        assert msgs[2] == "⛔ @C1 — ограничен бессрочно"  # no contacts: no tail

    def test_check_result_is_kept_for_the_accounts_list(self, monkeypatch, tmp_path):
        from modules import restricted_workers

        monkeypatch.chdir(tmp_path)
        s = self._S("sessions/a.jsession", reply="Unfortunately...\nyou are limited until 12 Nov 2026.")
        storage = self._storage(tmp_path, [s])

        self._run(storage)
        assert restricted_workers.load_status() == {"sessions/a.jsession": "12 Nov 2026"}
        assert storage.sessions == [s]  # stays in sessions/

        s.me = RuntimeError("network")  # a failed check keeps the entry
        self._run(storage)
        assert restricted_workers.load_status() == {"sessions/a.jsession": "12 Nov 2026"}

        s.me, s.reply = True, "Unfortunately...\nyou are limited forever."  # permanent: no date
        self._run(storage)
        assert restricted_workers.load_status() == {}

        s.reply = "Unfortunately...\nyou are limited until 12 Nov 2026."
        self._run(storage)
        s.reply = "Good news, no limits are currently applied."
        self._run(storage)
        assert restricted_workers.load_status() == {"sessions/a.jsession": "active"}


class TestPmMailingPause:
    def test_reversed_range(self, monkeypatch):
        from functions.pmmailing import PmMailingFunc

        monkeypatch.setattr("functions.pmmailing.asyncio.sleep", _no_sleep)
        fn = PmMailingFunc(_Storage(), ns(delay=[0], account_pause=[60, 30]))
        msgs, fn._report = collect()
        fn._active_accounts = {"first"}

        asyncio.run(fn.pause_between_accounts("second", "B"))  # randint(60, 30) used to raise

        assert any("пауза" in m for m in msgs)


class TestClearChatsOffset:
    def test_repeats_until_offset_zero(self):
        from functions.clear_chats import ClearDialogsFunc

        offsets = iter([100, 0])
        calls = []

        class _S:
            async def get_me(self):
                return ns(first_name="w")

            async def iter_dialogs(self):
                yield ns(entity=ns(), id=1, title="chat")

            async def __call__(self, request):
                calls.append(request)
                return ns(offset=next(offsets))

        fn = ClearDialogsFunc(_Storage(), ns(delay=[0]))
        _, report = collect()
        asyncio.run(fn.clear(_S(), report))

        assert len(calls) == 2


class TestParseMessageLink:
    def test_query_string(self):
        assert BaseFunction.parse_message_link("https://t.me/durov/123?single") == ("durov", 123)

    def test_public_topic(self):
        assert BaseFunction.parse_message_link("https://t.me/durov/5/123") == ("durov", 123)

    def test_private_topic(self):
        peer, message_id = BaseFunction.parse_message_link("https://t.me/c/1234567890/5/89?comment=1")
        assert (peer.channel_id, message_id) == (1234567890, 89)


class TestCommentLinks:
    """A …/<post>?comment=<id> link: the reaction / vote lands on the comment, in the discussion group."""

    class _S:
        def __init__(self, linked=50):
            self.linked, self.requests = linked, []

        async def get_me(self):
            return ns(first_name="A")

        async def __call__(self, request):
            from telethon.tl.functions.channels import GetFullChannelRequest

            if isinstance(request, GetFullChannelRequest):
                return ns(full_chat=ns(linked_chat_id=self.linked),
                          chats=[ns(id=10, title="channel"), ns(id=50, title="discussion")])
            self.requests.append(request)

    def test_comment_id(self):
        assert BaseFunction.comment_id("https://t.me/durov/123?comment=45") == 45
        assert BaseFunction.comment_id("https://t.me/durov/123?single&comment=45") == 45
        assert BaseFunction.comment_id("https://t.me/durov/123") is None

    def test_reaction_goes_on_the_comment(self):
        from functions.reactions import ReactionsFunc

        session = self._S()
        fn = ReactionsFunc(_Storage(), ns(delay=[0]))
        _, report = collect()
        asyncio.run(fn.set_reaction(session, "durov", 123, report, reaction="👍", comment=45))
        (request,) = session.requests
        assert (request.peer.id, request.msg_id) == (50, 45)

    def test_no_discussion_group_is_reported(self):
        from functions.reactions import ReactionsFunc

        session = self._S(linked=None)
        fn = ReactionsFunc(_Storage(), ns(delay=[0]))
        msgs, report = collect()
        asyncio.run(fn.set_reaction(session, "durov", 123, report, reaction="👍", comment=45))
        assert session.requests == [] and "нет чата обсуждения" in msgs[-1]

    def test_without_a_comment_the_post_itself(self):
        from functions.reactions import ReactionsFunc

        session = self._S()
        fn = ReactionsFunc(_Storage(), ns(delay=[0]))
        _, report = collect()
        asyncio.run(fn.set_reaction(session, "durov", 123, report, reaction="👍"))
        (request,) = session.requests
        assert (request.peer, request.msg_id) == ("durov", 123)


class TestReportPeer:
    """Telethon can't resolve a post link as a peer: the report takes the post's chat from it."""

    @staticmethod
    def _peer(link):
        from functions.report import ReportFunc

        sent = []

        class _Session:
            async def __call__(self, request):
                sent.append(request.peer)

            async def get_input_entity(self, peer):  # a private link's chat, from the cache
                return peer

        asyncio.run(ReportFunc(_Storage(), ns(delay=[0])).report_step(_Session(), link, [1], "", b""))
        return sent[0]

    def test_a_post_link_reports_in_its_chat(self):
        assert self._peer("https://t.me/chan/123") == "chan"
        assert self._peer("t.me/c/123/45").channel_id == 123

    def test_a_private_chat_link_reports_in_that_chat(self):
        """t.me/c/<id> alone (the post ids are asked separately) is the chat, not a chat named "c"."""
        for link in ("https://t.me/c/123", "t.me/c/123/", "https://t.me/c/123?single"):
            assert self._peer(link).channel_id == 123

    def test_a_chat_link_or_name_goes_as_is(self):
        for peer in ("https://t.me/chan", "@chan"):
            assert self._peer(peer) == peer

    def test_a_numeric_id_is_the_chat(self):
        """Telethon reads a "-100…" string as a phone number: the id is turned into the chat's peer."""
        from telethon import types

        assert self._peer("-1001234567890") == types.PeerChannel(1234567890)
        assert self._peer("-123") == types.PeerChat(123)


# --- modules/ ----------------------------------------------------------------

class TestSetupBroadcast:
    def test_requires_one_message(self, monkeypatch):
        from modules.settings import Settings

        answers = iter(["", "hi", "", "1-3", "go"])
        monkeypatch.setattr("modules.settings.console.input", lambda *_: next(answers))
        assert Settings.setup_broadcast() == (["hi"], [1, 3], "go")

    def test_reasks_a_delay_of_more_than_two_parts(self, monkeypatch):
        from modules.settings import Settings

        # "1-2-3" would be written to config.toml and then rejected by Settings() on the next start
        answers = iter(["hi", "", "1-2-3", "2-5", "go"])
        monkeypatch.setattr("modules.settings.console.input", lambda *_: next(answers))
        assert Settings.setup_broadcast() == (["hi"], [2, 5], "go")


class TestUpdaterPullFailure:
    def test_git_command_error_is_reported(self, monkeypatch):
        from modules import updater

        class _Remote:
            def pull(self):
                raise GitCommandError("pull", 1)

        monkeypatch.setattr(updater, "Repo",
                            lambda *_: ns(remote=lambda name: _Remote(), head=ns(commit=None)))
        printed = []
        console = ns(status=lambda *_: contextlib.nullcontext(), print=printed.append)

        updater.update(console)  # must not crash main before the menu

        assert any("Обновление не удалось" in p for p in printed)


# --- bot/ --------------------------------------------------------------------

class _Manager:
    def __init__(self):
        self.released = 0

    def release(self):
        self.released += 1


def _callback(value):
    answers = []

    async def answer(*a, **k):
        answers.append(a)

    cb = ns(message=ns(chat=ns(id=7), answer=answer), answer=answer, bot=None)
    return cb, ns(value=str(value)), answers


class TestReportFlowConcurrency:
    def teardown_method(self):
        from bot.routers import moderation
        moderation._FLOWS.clear()

    def test_double_tap_does_not_run_second_step(self):
        from bot.routers import moderation

        stepped = []

        async def step(flow, option):
            stepped.append(option)
            return "choose", []

        flow = {"busy": True, "selections": []}
        moderation._FLOWS[7] = (ns(step=step), flow, [ns(option=b"a")])
        cb, data, _ = _callback(0)
        asyncio.run(moderation.rm_choose(cb, data, _Manager(), ns(in_job=[])))

        assert stepped == [] and flow["selections"] == []

    def test_abort_during_step_does_not_release_foreign_slot(self):
        from bot.routers import moderation

        async def step(flow, option):
            moderation._FLOWS.pop(7, None)  # /cancel (_abort) ran while we awaited
            raise RuntimeError("disconnected")

        class _Session:
            async def disconnect(self):
                return None

        flow = {"busy": False, "selections": [], "session": _Session()}
        moderation._FLOWS[7] = (ns(step=step), flow, [ns(option=b"a")])
        cb, data, _ = _callback(0)
        manager = _Manager()
        asyncio.run(moderation.rm_choose(cb, data, manager, ns(in_job=[])))

        assert manager.released == 0
        assert 7 not in moderation._FLOWS

    def test_stale_index_ignored(self):
        from bot.routers import moderation

        flow = {"busy": False, "selections": []}
        moderation._FLOWS[7] = (ns(step=None), flow, [ns(option=b"a")])
        cb, data, _ = _callback(5)
        asyncio.run(moderation.rm_choose(cb, data, _Manager(), ns(in_job=[])))

        assert flow["selections"] == [] and flow["busy"] is False


class TestReporterLength:
    def test_text_capped_at_4096(self):
        from bot.services.runner import TelegramReporter

        sent = []

        class _Bot:
            async def edit_message_text(self, text, **k):
                sent.append(text)

        reporter = TelegramReporter(_Bot(), 1, header="H")
        reporter.message_id = 1
        reporter.lines = ["x" * 500] * 25
        asyncio.run(reporter.finish("Готово"))

        assert sent and len(sent[-1]) <= 4096
        assert sent[-1].startswith("Готово")


class TestMailLimit:
    def _run(self, text):
        from bot.routers import broadcasts

        replies, data, states = [], {}, []

        class _State:
            async def update_data(self, **kw):
                data.update(kw)

            async def set_state(self, s):
                states.append(s)

        async def answer(t, **k):
            replies.append(t)

        asyncio.run(broadcasts.mail_limit(ns(text=text, answer=answer), _State()))
        return replies, data, states

    def test_typo_reprompts(self):
        replies, data, states = self._run("1O0")
        assert data == {} and states == [] and "положительное" in replies[0]

    def test_dash_means_all(self):
        _, data, states = self._run("-")
        assert data == {"limit": None} and states

    def test_number(self):
        _, data, _ = self._run("100")
        assert data == {"limit": 100}


class TestReadPreview:
    def test_cut_never_splits_an_entity(self):
        from bot.routers.scraping import read_preview

        df = pd.DataFrame({"c": ["&" * 400] * 10})
        out = read_preview(df)
        body = out[len("<pre>"):out.index("</pre>")]
        assert html.escape(html.unescape(body)) == body  # well-formed escaping


# --- scraper/ ----------------------------------------------------------------

class TestExcelFormulas:
    def test_equals_prefixed_text_stays_text(self, tmp_path):
        from scraper.datafiles import save_table

        path = save_table(pd.DataFrame({"Content": ["=== NEWS ===", "ok"]}), tmp_path / "x", "excel")
        cell = openpyxl.load_workbook(path)["Sheet1"]["A2"]

        assert (cell.value, cell.data_type) == ("=== NEWS ===", "s")


class TestJoinReturnsTarget:
    def _join(self, link, mode, reply):
        from functions.joiner import JoinerFunc

        class _S:
            async def __call__(self, request):
                return reply(request)

        fn = JoinerFunc(_Storage(), ns(delay=[0]))
        _, report = collect()
        return asyncio.run(fn.join(_S(), link, 0, mode, report))

    def test_invite_link_returns_joined_chat(self):
        chat = ns(id=5)
        assert self._join("https://t.me/joinchat/abc", "1", lambda r: ns(chats=[chat])) is chat

    def test_falls_back_to_link_without_chats(self):
        assert self._join("@chan", "1", lambda r: ns(chats=[])) == "@chan"

    def test_mode2_returns_linked_chat(self):
        from telethon.tl.functions.channels import GetFullChannelRequest

        linked = ns(id=9)

        def reply(request):
            if isinstance(request, GetFullChannelRequest):
                return ns(full_chat=ns(linked_chat_id=9), chats=[ns(id=1), linked])
            return ns(chats=[])

        assert self._join("@channel", "2", reply) is linked

    def test_webview_join_is_not_counted(self):
        # the join waits for a confirmation in the chat's bot web app: not joined
        from telethon.tl.types.messages import ChatInviteJoinResultWebView

        webview = ChatInviteJoinResultWebView(bot_id=1, query_id=2, users=[])
        assert self._join("@chan", "1", lambda r: webview) is False


class TestAddContactsUsernameFallback:
    def _add(self, row):
        from telethon.tl.types import InputUser
        from functions.add_contacts import AddContactsFunc

        added_ids = []

        class _S:
            async def __call__(self, request):
                if isinstance(request.id, InputUser):  # foreign access_hash
                    raise ValueError("USER_ID_INVALID")
                added_ids.append(request.id)

            async def get_input_entity(self, username):
                return ns(username=username)

            async def get_me(self):
                return ns(id=7)

        fn = AddContactsFunc(_Storage(), ns(delay=[0]))
        msgs, fn._report = collect()
        fn.added, fn.ledger, fn._unsaved, fn._my_ids = 0, {}, 0, {}
        fn._limits = ns(reached=lambda k: False, bump=lambda k: None, cap=0)
        asyncio.run(fn.add_one(_S(), row))
        return fn.added, added_ids, msgs

    def test_falls_back_to_username(self):
        added, ids, _ = self._add({"user_id": 1, "access_hash": 2, "username": "bob"})
        assert added == 1 and ids[0].username == "bob"

    def test_no_username_is_skipped(self):
        added, _, msgs = self._add({"user_id": 1, "access_hash": 2})
        assert added == 0 and any("пропуск user_id=1" in m for m in msgs)


class TestAddContactsName:
    def _first_name(self, row):
        from functions.add_contacts import AddContactsFunc

        sent = []

        class _S:
            async def __call__(self, request):
                sent.append(request)

            async def get_me(self):
                return ns(id=7)

        fn = AddContactsFunc(_Storage(), ns(delay=[0]))
        msgs, fn._report = collect()
        fn.added, fn.ledger, fn._unsaved, fn._my_ids = 0, {}, 0, {}
        fn._limits = ns(reached=lambda k: False, bump=lambda k: None, cap=0)
        asyncio.run(fn.add_one(_S(), {"user_id": 1, "access_hash": 2, **row}))
        return sent[0].first_name

    def test_scrape_base_full_name(self):
        assert self._first_name({"name": "Иван Петров"}) == "Иван Петров"

    def test_first_name_column_wins(self):
        assert self._first_name({"first_name": "Ivan", "name": "Иван Петров"}) == "Ivan"

    def test_username_when_no_name(self):
        assert self._first_name({"name": "", "username": "bob"}) == "bob"

    def test_placeholder_when_nothing(self):
        assert self._first_name({}) == "contact"


class TestAddContactsDailyCap:
    def test_second_add_is_capped(self, tmp_path):
        from functions.add_contacts import AddContactsFunc
        from functions.base.base import AccountLimited
        from modules.account_limits import DailyCounter

        class _S:
            async def get_me(self):
                return ns(id=7)

            async def __call__(self, request):
                pass  # AddContactRequest succeeds

        fn = AddContactsFunc(_Storage(), ns(delay=[0]))
        _msgs, fn._report = collect()
        fn.added, fn.ledger, fn._unsaved, fn._my_ids = 0, {}, 0, {}
        fn._limits = DailyCounter(str(tmp_path / "contacts_limits.json"), 1)

        session, row = _S(), {"user_id": 1, "access_hash": 2}
        asyncio.run(fn.add_one(session, row))  # first: added + bump
        assert fn.added == 1

        with pytest.raises(AccountLimited):
            asyncio.run(fn.add_one(session, {"user_id": 2, "access_hash": 3}))  # cap 1 reached


class TestProfilePhotoSkipsHiddenFiles:
    def test_only_regular_visible_files(self, monkeypatch, tmp_path):
        from functions.change_profile_photo import ChangeProfilePhotoFunc

        photos = tmp_path / "assets" / "photos"
        (photos / "sub").mkdir(parents=True)
        (photos / ".DS_Store").write_bytes(b"\x00")
        (photos / "a.jpg").write_bytes(b"jpg")
        monkeypatch.chdir(tmp_path)

        fn = ChangeProfilePhotoFunc(_Storage([object()] * 5), ns(delay=[0], profile_pause=[0]))
        picked = []

        async def set_profile_photo(session, path, report):
            picked.append(path)

        fn.set_profile_photo = set_profile_photo
        _, report = collect()
        asyncio.run(fn.run(report))

        assert picked and all(p.endswith("a.jpg") for p in picked)


class TestCancelNonCancelable:
    def test_reports_job_keeps_running(self):
        from bot.routers import menu
        from bot.services.jobs import JobManager

        replies = []

        async def answer(text, **k):
            replies.append(text)

        class _State:
            async def get_state(self):
                return None

            async def clear(self):
                return None

        async def scenario():
            manager = JobManager()
            manager.acquire("Скрапинг", cancelable=False, timeout=0)
            await menu.cancel(ns(answer=answer), _State(), manager)
            return manager

        manager = asyncio.run(scenario())
        assert manager.active and "не прерывается" in replies[0]


class TestMailRunStats:
    def _run(self, active_after_load, load_stats):
        from bot.routers import broadcasts

        replies = []

        async def answer(text, **k):
            replies.append(text)

        class _State:
            async def get_data(self):
                return {"path": "x", "skip": True}

            async def clear(self):
                return None

        class _Manager:
            label = "Рассылка"
            checks = 0

            @property
            def active(self):
                type(self).checks += 1
                return active_after_load and type(self).checks > 1  # busy only after the await

        instance = ns(load_recipients=lambda p: ["a"], load_stats=load_stats,
                      filter_unsent=lambda r: r)
        functions = {"PmMailingFunc": instance}
        asyncio.run(broadcasts.mail_run(ns(answer=answer), _State(), None, None, functions,
                                        _Manager(), None))
        return replies

    def test_broken_stats_is_reported(self):
        def broken():
            raise ValueError("bad json")

        replies = self._run(False, broken)
        assert any("Статистика не прочитана" in r for r in replies)

    def test_job_started_meanwhile_keeps_its_stats(self):
        loaded = []
        replies = self._run(True, lambda: loaded.append(1))
        assert loaded == [] and any("Занят" in r for r in replies)


# --- round 3 -----------------------------------------------------------------

class TestCliLoopReuse:
    def test_functions_run_on_the_sessions_loop(self, tmp_path):
        from modules.storages.functions_storage import FunctionsStorage

        loop = asyncio.new_event_loop()
        try:
            fs = FunctionsStorage(str(tmp_path), ns(loop=loop, sessions=[]), ns())
            assert fs.loop is loop  # Telethon refuses calls from another loop
        finally:
            asyncio.set_event_loop(None)
            loop.close()


class TestStickerRequest:
    def test_hash_is_passed(self, monkeypatch):
        from functions.broadcast import Broadcast

        requests = []

        class _S:
            async def __call__(self, request):
                requests.append(request)
                return ns(documents=["doc"])

            async def send_file(self, *a, **k):
                return None

        fn = Broadcast(_Storage(), ns(delay=[0]))
        fn.configure(choice=4, sticker_set="https://t.me/addstickers/Pack", content=None)

        async def no_send(*a, **k):
            return None

        import functions.broadcast as mod
        monkeypatch.setattr(mod.rich_message, "send", no_send)
        asyncio.run(fn._send(_S(), "chat", None, None))

        assert requests[0].hash == 0 and requests[0].stickerset.short_name == "Pack"


class TestUserBannedInChannel:
    def test_is_account_limited_not_a_chat_leave(self):
        from telethon.errors import UserBannedInChannelError
        from functions.base.base import AccountLimited

        async def make():
            raise UserBannedInChannelError(request=None)

        fn = BaseFunction()
        try:
            asyncio.run(fn.safe_call(make))
            assert False, "expected AccountLimited"
        except AccountLimited:
            pass


class TestJobManagerStops:
    def test_double_stop_still_frees_slot(self, monkeypatch):
        from bot.services import jobs as jobs_mod
        from bot.services.jobs import JobManager

        class _Reporter:
            async def start(self):
                return None

            async def finish(self, text):
                await asyncio.sleep(0)

        class _Session:
            async def disconnect(self):
                await asyncio.sleep(0.01)

        monkeypatch.setattr(jobs_mod, "TelegramReporter", lambda *a, **k: _Reporter())

        async def scenario():
            manager = JobManager()

            async def factory(f, reporter):
                await asyncio.sleep(10)

            pool = ns(workers=[_Session()], run=lambda inst, bf, fac, rep: fac(None))
            await manager.run(None, 1, pool, None, None, factory, "job", "done")
            await asyncio.sleep(0)
            await manager.stop()
            await asyncio.sleep(0.005)  # inside _wrap's cleanup (disconnect)
            await manager.stop()        # second tap
            await asyncio.sleep(0.05)
            return manager

        manager = asyncio.run(scenario())
        assert manager.active is False

    def test_stop_while_starting_aborts_start(self, monkeypatch):
        from bot.services import jobs as jobs_mod
        from bot.services.jobs import JobManager

        manager = JobManager()

        class _Reporter:
            async def start(self):
                await manager.stop()  # /cancel lands while the status message is posted

            async def finish(self, text):
                return None

        monkeypatch.setattr(jobs_mod, "TelegramReporter", lambda *a, **k: _Reporter())
        started = asyncio.run(manager.run(None, 1, ns(workers=[]), None, None,
                                          lambda f, r: None, "job", "done"))

        assert started is False and manager.active is False


class TestReportCancelDuringConnect:
    def test_no_step_after_abort(self):
        from bot.routers import moderation

        stepped = []

        class _First:
            async def connect(self):
                moderation._FLOWS.pop(7, None)  # /cancel ran _abort while connecting

            async def disconnect(self):
                return None

        async def step(flow, option):
            stepped.append(option)
            return "done", None

        class _Manager:
            label = ""

            def acquire(self, *a, **k):
                return True

            def release(self):
                raise AssertionError("slot no longer ours")

        class _State:
            async def get_data(self):
                return {"link": "l", "ids": [1]}

            async def clear(self):
                return None

        async def answer(*a, **k):
            return None

        instance = ns(step=step)
        pool = ns(delegate=lambda inst: [_First()])
        functions = {"ReportFunc": instance}
        msg = ns(text="-", chat=ns(id=7), answer=answer, bot=None)
        try:
            asyncio.run(moderation.rm_begin(msg, _State(), pool, functions, _Manager()))
        finally:
            moderation._FLOWS.clear()

        assert stepped == []


class TestCancelInsideForm:
    def test_form_cancel_keeps_job(self):
        from bot.routers import menu

        stopped = []
        replies = []

        class _State:
            async def get_state(self):
                return "ChangeBio:text"

            async def clear(self):
                return None

        async def stop():
            stopped.append(1)
            return True

        async def answer(text, **k):
            replies.append(text)

        asyncio.run(menu.cancel(ns(answer=answer, chat=ns(id=1)), _State(), ns(stop=stop, active=True, label="Рассылка")))
        assert stopped == [] and replies[0].startswith("Отменено.") and "/cancel ещё раз" in replies[0]


class TestReporterTrailingFlush:
    def test_burst_tail_is_shown(self):
        from bot.services.runner import TelegramReporter

        sent = []

        class _Bot:
            async def edit_message_text(self, text, **k):
                sent.append(text)

        async def scenario():
            r = TelegramReporter(_Bot(), 1, header="H", min_interval=0.05)
            r.message_id = 1
            await r("first")
            await r("second")  # inside min_interval
            await asyncio.sleep(0.1)

        asyncio.run(scenario())
        assert "second" in sent[-1]


class TestVideoSticker:
    def test_saved_as_webm(self):
        from bot.services.capture import _source_and_flags

        sticker = ns(is_video=True, is_animated=False)
        msg = ns(photo=None, video=None, animation=None, voice=None, video_note=None,
                 audio=None, sticker=sticker, document=None)
        assert _source_and_flags(msg)[1] == "sticker.webm"


class TestMailLimitSuperscript:
    def test_superscript_digit_reprompts(self):
        replies, data, _ = TestMailLimit()._run("²")
        assert data == {} and "положительное" in replies[0]


class TestNewFilesOverwritten:
    def test_rewritten_file_is_new(self, tmp_path):
        import os
        from bot.services.scraping import dir_snapshot, new_files

        f = tmp_path / "x_posts.parquet"
        f.write_text("old")
        before = dir_snapshot(str(tmp_path))
        os.utime(f, ns=(1, before["x_posts.parquet"] + 10**9))  # overwritten later
        assert new_files(str(tmp_path), before) == [str(f)]


class TestExcelHeaderFormula:
    def test_header_stays_text(self, tmp_path):
        from scraper.datafiles import save_table

        path = save_table(pd.DataFrame({"=promo": [1]}), tmp_path / "h", "excel")
        cell = openpyxl.load_workbook(path)["Sheet1"]["A1"]
        assert (cell.value, cell.data_type) == ("=promo", "s")


class TestScrapeFuncSwappedDates:
    def test_reports_and_does_not_run(self, monkeypatch):
        from functions import scraper as fs

        answers = iter(["name", "out", "@chan", "31.12.2024", "01.01.2024", ""])
        monkeypatch.setattr(fs, "pick_session",
                            lambda storage, personal=None: ns(path="sessions/a.jsession", client=object(),
                                                              personal=False))
        monkeypatch.setattr(fs.Prompt, "ask", lambda *a, **k: next(answers))
        monkeypatch.setattr(fs.Confirm, "ask", lambda *a, **k: False)
        printed = []
        monkeypatch.setattr(fs.console, "print", lambda *a, **k: printed.append(str(a[0])))
        ran = []
        monkeypatch.setattr(fs, "scrape_run", lambda *a: ran.append(1))

        f = fs.ScrapeFunc(_Storage(), ns(delay=[0], api_id=1, api_hash="h"))
        f.ask_int = lambda *a, **k: 100
        f.execute()

        assert ran == [] and any("позже даты конца" in p for p in printed)


class TestInvitingMissingInvitees:
    def test_privacy_restricted_not_counted(self):
        from functions.inviting import InvitingFunc

        users = [ns(id=1, bot=False, deleted=False, is_self=False),
                 ns(id=2, bot=False, deleted=False, is_self=False)]

        class _S:
            async def get_me(self):
                return ns(first_name="A")

            async def get_entity(self, dest):
                return dest

            async def iter_participants(self, source):
                for u in users:
                    yield u

            async def __call__(self, request):
                return ns(missing_invitees=[ns(user_id=1)] if request.users[0].id == 1 else [])

        fn = InvitingFunc(_Storage(), ns(delay=[0]))
        fn.delay = _no_sleep
        fn._limits = ns(reached=lambda k: False, bump=lambda k: None, cap=0)

        async def resolve_source(session, link):
            return "src"

        fn.resolve_source = resolve_source
        msgs, report = collect()
        asyncio.run(fn.invite(_S(), "src", "dest", [1, 2], report))

        invited = [m for m in msgs if "приглашён" in m]
        assert invited == ["[A] приглашён 2, всего: 1"]

    def test_delay_after_privacy_error(self):
        # the invite request was sent even when it errors: the delay must still apply
        from telethon.errors import UserPrivacyRestrictedError
        from functions.inviting import InvitingFunc

        users = [ns(id=1, bot=False, deleted=False, is_self=False),
                 ns(id=2, bot=False, deleted=False, is_self=False)]

        class _S:
            async def get_me(self):
                return ns(first_name="A")

            async def get_entity(self, dest):
                return dest

            async def iter_participants(self, source):
                for u in users:
                    yield u

            async def __call__(self, request):
                raise UserPrivacyRestrictedError(request=None)

        fn = InvitingFunc(_Storage(), ns(delay=[0]))
        fn._limits = ns(reached=lambda k: False, bump=lambda k: None, cap=0)
        delays = []

        async def count_delay():
            delays.append(1)

        fn.delay = count_delay

        async def resolve_source(session, link):
            return "src"

        fn.resolve_source = resolve_source
        _msgs, report = collect()
        asyncio.run(fn.invite(_S(), "src", "dest", [1, 2], report))

        assert len(delays) == 2


# --- round 4 -----------------------------------------------------------------

class TestGifSentAsAnimation:
    def test_capture_marks_animation(self):
        from bot.services.capture import _source_and_flags

        msg = ns(photo=None, video=None, animation=ns(file_name=None), voice=None,
                 video_note=None, audio=None, sticker=None, document=None)
        found = _source_and_flags(msg)
        assert found[1] == "animation.mp4" and found[-1] is True

    def test_send_passes_animated_attribute(self):
        from telethon import types as tl
        from modules.rich_message import MediaItem, RichContent, _send_once

        calls = []

        class _S:
            async def send_file(self, peer, path, **kw):
                calls.append(kw)

        async def safe_call(make):
            return await make()

        content = RichContent(text="", media=[MediaItem(path="a.mp4", animated=True)])
        asyncio.run(_send_once(_S(), "peer", content, [], safe_call))
        assert isinstance(calls[0]["attributes"][0], tl.DocumentAttributeAnimated)


class TestSessionFiles:
    def test_ipv6_string_session_loaded_garbage_skipped(self, tmp_path):
        from telethon.sessions import StringSession
        from modules.storages.sessions_storage import SessionsStorage

        ss = StringSession()
        ss.set_dc(2, "2001:67c:4e8:f002::a", 443)
        ss._auth_key = __import__("telethon.crypto", fromlist=["AuthKey"]).AuthKey(b"\x01" * 256)
        v6 = ss.save()
        assert len(v6) == 369
        (tmp_path / "v6.session").write_text(v6)
        (tmp_path / "junk.session").write_text("x" * 100)

        storage = SessionsStorage(str(tmp_path), 1, "hash", initialize=False)
        assert list(storage.full_sessions) == [str(tmp_path / "v6.session")]

    def test_string_session_has_entities_set(self):
        # the bot clears this private set after each job / autoreply round to bound memory;
        # a Telethon upgrade that renames it would silently turn that clearing into a no-op
        from telethon.sessions import StringSession

        assert isinstance(StringSession()._entities, set)

    def test_broken_jsession_skipped(self, tmp_path):
        from modules.storages.sessions_storage import SessionsStorage

        (tmp_path / "broken.jsession").write_text("{not json")
        storage = SessionsStorage(str(tmp_path), 1, "hash", initialize=False)
        assert storage.sessions == []


class TestUpdaterAheadOfOrigin:
    def test_no_update_when_origin_is_ancestor(self, monkeypatch):
        from modules import updater

        upcoming = ns(hexsha="origin", message="m")
        repo = ns(
            remotes=ns(origin=ns(refs=ns(master=ns(commit=upcoming)))),
            head=ns(commit=ns(hexsha="local")),
            is_ancestor=lambda a, b: a is upcoming,
        )
        monkeypatch.setattr(updater.git, "Repo", lambda *a: repo)
        monkeypatch.setattr(updater.git, "Remote", lambda r, name: ns(fetch=lambda: ["info"]))
        monkeypatch.setattr(updater, "get_current_commit", lambda: "local")

        assert updater.check_update() == {"has_update": False}


# --- round 5 -----------------------------------------------------------------

class TestInviteJoinResult:
    def test_inviting_reads_chat_from_wrapped_updates(self):
        from telethon.tl.functions.messages import ImportChatInviteRequest
        from functions.inviting import InvitingFunc

        chat = ns(id=3)

        class _S:
            async def __call__(self, request):
                assert isinstance(request, ImportChatInviteRequest)
                return ns(updates=ns(chats=[chat]))  # ChatInviteJoinResultOk

        fn = InvitingFunc(_Storage(), ns(delay=[0]))
        assert asyncio.run(fn.resolve_source(_S(), "https://t.me/+abcdef")) is chat

    def test_inviting_webview_join_is_refused(self):
        from telethon.tl.types.messages import ChatInviteJoinResultWebView
        from functions.inviting import InvitingFunc

        class _S:
            async def __call__(self, request):
                return ChatInviteJoinResultWebView(bot_id=1, query_id=2, users=[])

        fn = InvitingFunc(_Storage(), ns(delay=[0]))
        with pytest.raises(ValueError, match="web-app"):
            asyncio.run(fn.resolve_source(_S(), "https://t.me/+abcdef"))

    def test_joiner_reads_chat_from_wrapped_updates(self):
        chat = ns(id=5)
        got = TestJoinReturnsTarget()._join("https://t.me/joinchat/abc", "1",
                                            lambda r: ns(updates=ns(chats=[chat])))
        assert got is chat


class TestMentionsAndCaptions:
    def test_mention_converted_to_input_entity(self):
        from telethon import types as tl
        from modules.rich_message import RichContent, send

        sent = []

        class _S:
            async def get_input_entity(self, user_id):
                if user_id == 2:
                    raise ValueError("unknown")
                return tl.InputUser(user_id, 99)

            async def send_message(self, peer, text, formatting_entities=None, **kw):
                sent.append(formatting_entities)

        async def safe_call(make):
            return await make()

        entities = [tl.MessageEntityMentionName(0, 2, user_id=1),
                    tl.MessageEntityMentionName(2, 2, user_id=2),
                    tl.MessageEntityBold(0, 4)]
        asyncio.run(send(_S(), "peer", RichContent(text="abcd", entities=entities), safe_call))

        kinds = [type(e).__name__ for e in sent[0]]
        assert kinds == ["InputMessageEntityMentionName", "MessageEntityBold"]  # unknown user dropped

    def test_caption_not_parsed_as_markdown(self):
        from modules.rich_message import MediaItem, RichContent, _send_once

        calls = []

        class _S:
            async def send_file(self, peer, path, **kw):
                calls.append(kw)

        async def safe_call(make):
            return await make()

        content = RichContent(text="use **kwargs", media=[MediaItem(path="a.jpg")])
        asyncio.run(_send_once(_S(), "peer", content, [], safe_call))
        assert calls[0]["parse_mode"] is None


class TestClearChatsChannelForbidden:
    def test_banned_channel_is_left(self):
        from telethon import types as tl
        from telethon.tl.functions.channels import LeaveChannelRequest
        from functions.clear_chats import ClearDialogsFunc

        calls = []
        forbidden = tl.ChannelForbidden(id=1, access_hash=2, title="x")

        class _S:
            async def get_me(self):
                return ns(first_name="w")

            async def iter_dialogs(self):
                yield ns(entity=forbidden, id=-1001, title="x")

            async def __call__(self, request):
                calls.append(request)

        fn = ClearDialogsFunc(_Storage(), ns(delay=[0]))
        _, report = collect()
        asyncio.run(fn.clear(_S(), report))
        assert isinstance(calls[0], LeaveChannelRequest)


class TestScraperCredentials:
    def test_account_proxy_and_app_carried_over(self):
        from modules.scraper_creds import build_credentials

        init = ns(device_model="Redmi Note 10", system_version="SDK 31", app_version="10.3",
                  lang_code="en", system_lang_code="en-US")
        client = ns(api_id=7, api_hash="app", session=ns(save=lambda: "KEY"),
                    _proxy=("socks5", "1.1.1.1", 1080), _init_request=init)
        creds = build_credentials(client)
        assert (creds.api_id, creds.api_hash, creds.session_string, creds.proxy) == \
            (7, "app", "KEY", ("socks5", "1.1.1.1", 1080))
        assert creds.device == {"device_model": "Redmi Note 10", "system_version": "SDK 31",
                                "app_version": "10.3", "lang_code": "en", "system_lang_code": "en-US"}


class TestVerifyFuncSwappedDates:
    def test_reports_and_does_not_run(self, monkeypatch):
        from functions import scraper as fs

        answers = iter(["posts.parquet", "@chan", "31.12.2024", "01.01.2024", ""])
        monkeypatch.setattr(fs, "pick_session", lambda storage, personal=None: object())
        monkeypatch.setattr(fs.Prompt, "ask", lambda *a, **k: next(answers))
        printed = []
        monkeypatch.setattr(fs.console, "print", lambda *a, **k: printed.append(str(a[0])))
        ran = []
        monkeypatch.setattr(fs, "verify_run", lambda *a: ran.append(1))

        f = fs.VerifyFunc(_Storage(), ns(delay=[0], api_id=1, api_hash="h"))
        f.ask_int = lambda *a, **k: 0
        f.execute()

        assert ran == [] and any("позже даты конца" in p for p in printed)


class TestExcelErrorStrings:
    def test_error_code_text_stays_text(self, tmp_path):
        from scraper.datafiles import save_table

        path = save_table(pd.DataFrame({"Content": ["#N/A"]}), tmp_path / "e", "excel")
        cell = openpyxl.load_workbook(path)["Sheet1"]["A2"]
        assert (cell.value, cell.data_type) == ("#N/A", "s")


class TestJobSlotBackstop:
    def test_cancel_before_first_step_frees_slot(self, monkeypatch):
        from bot.services import jobs as jobs_mod
        from bot.services.jobs import JobManager

        class _Reporter:
            async def start(self):
                return None

            async def finish(self, text):
                return None

        async def scenario():
            manager = JobManager()
            await manager.run(None, 1, ns(workers=[]), None, None, lambda f, r: None, "job", "done")
            manager._task.cancel()  # before _wrap ever runs: its try/finally never starts
            await asyncio.sleep(0.01)
            return manager

        monkeypatch.setattr(jobs_mod, "TelegramReporter", lambda *a, **k: _Reporter())
        manager = asyncio.run(scenario())
        assert manager.active is False


class TestReportAbortDuringAnswer:
    def test_no_step(self):
        from bot.routers import moderation

        stepped = []

        async def step(flow, option):
            stepped.append(option)
            return "choose", []

        async def answer(*a, **k):
            moderation._FLOWS.pop(7, None)  # /cancel ran _abort meanwhile

        flow = {"busy": False, "selections": []}
        moderation._FLOWS[7] = (ns(step=step), flow, [ns(option=b"a")])
        cb = ns(message=ns(chat=ns(id=7), answer=answer), answer=answer, bot=None)
        try:
            asyncio.run(moderation.rm_choose(cb, ns(value="0"), _Manager(), ns(in_job=[])))
        finally:
            moderation._FLOWS.clear()
        assert stepped == []


class TestScrapeSnapshotBeforeSlot:
    def test_unreadable_out_dir_keeps_slot_free(self, monkeypatch):
        from bot.routers import scraping
        from bot.services import scraping as svc
        from bot.services.jobs import JobManager

        def boom(path):
            raise PermissionError("denied")

        monkeypatch.setattr(svc, "dir_snapshot", boom)

        replies = []

        class _Bot:
            async def send_message(self, chat_id, text, **k):
                replies.append(text)

        manager = JobManager()  # the scraper's slot
        pool = ns(storage=ns(sessions=[object()], jsessions_paths={},
                             get_session_path=lambda c: "sessions/w.jsession"))
        params = scraping._scrape_params({"out_dir": "/root/x", "channels": "@a", "name": "n",
                                          "date_min": "01.01.2024", "date_max": "02.01.2024",
                                          "account": "sessions/w.jsession"})
        asyncio.run(scraping._run_scrape(_Bot(), 1, manager, pool, None, params))

        assert manager.active is False and any("недоступна" in r for r in replies)


class TestCheckSessionDisconnects:
    def test_unauthorized_client_disconnected(self, tmp_path):
        from modules.storages.sessions_storage import SessionsStorage

        path = tmp_path / "a.session"
        path.write_text("x")
        storage = SessionsStorage(str(tmp_path), 1, "hash", initialize=False)

        class _S:
            disconnected = False

            async def connect(self):
                return None

            async def is_user_authorized(self):
                return False

            async def disconnect(self):
                self.disconnected = True

        s = _S()
        storage.full_sessions[str(path)] = s
        asyncio.run(storage.check_session(s, str(path)))
        assert s.disconnected and (tmp_path / "inactive" / "a.session").exists()


class TestUpdaterRequirements:
    def test_installs_when_requirements_changed(self, monkeypatch):
        from modules import updater

        new_head = object()
        old_head = ns(diff=lambda other: [ns(b_path="requirements.txt")] if other is new_head else [])
        repo = ns(head=ns(commit=old_head))

        def pull():
            repo.head.commit = new_head  # HEAD moves; FetchInfo would carry no old_commit

        repo.remote = lambda name: ns(pull=pull)
        monkeypatch.setattr(updater, "Repo", lambda *_: repo)
        installed, restarted = [], []
        monkeypatch.setattr(updater, "update_requirements", lambda console: installed.append(1))
        monkeypatch.setattr(updater, "restart_app", lambda: restarted.append(1))
        console = ns(status=lambda *_: contextlib.nullcontext(), print=lambda *a: None)

        updater.update(console)

        assert installed == [1] and restarted == [1]


# --- round 5 leftovers ---------------------------------------------------------

class TestReporterLimits:
    def _reporter(self, sent, fail_with=None):
        from bot.services.runner import TelegramReporter

        class _Bot:
            async def edit_message_text(self, text, **k):
                if fail_with is not None:
                    raise fail_with
                sent.append(text)

        r = TelegramReporter(_Bot(), 1, header="H")
        r.message_id = 1
        return r

    def test_utf16_and_huge_header_fit(self):
        sent = []
        r = self._reporter(sent)
        r.lines = ["😀" * 300] * 25  # 600 UTF-16 units per line
        asyncio.run(r.finish("⚠️ Ошибка: " + "x" * 5000))

        text = sent[-1]
        assert len(text.encode("utf-16-le")) // 2 <= 4096
        assert text.startswith("⚠️ Ошибка: xxx")

    def test_retry_after_pauses_progress_edits(self):
        from aiogram.exceptions import TelegramRetryAfter

        sent = []
        err = TelegramRetryAfter(method=None, message="flood", retry_after=30)
        r = self._reporter(sent, fail_with=err)

        async def scenario():
            await r("one")  # hits flood control
            return r._last

        import time
        last = asyncio.run(scenario())
        assert last > time.monotonic() + 20  # next progress edit only after the wait


class TestSendFilesErrors:
    def test_failed_upload_reported_and_rest_sent(self, tmp_path):
        from bot.routers.scraping import _send_files

        a, b = tmp_path / "a.parquet", tmp_path / "b.parquet"
        a.write_text("a")
        b.write_text("b")
        docs, msgs = [], []

        class _Bot:
            async def send_document(self, chat_id, doc):
                if doc.path.endswith("a.parquet"):
                    raise RuntimeError("timeout")
                docs.append(doc.path)

            async def send_message(self, chat_id, text):
                msgs.append(text)

        asyncio.run(_send_files(_Bot(), 1, [str(a), str(b)]))
        assert docs == [str(b)] and "Не удалось отправить a.parquet" in msgs[0]


class TestJobNotRunLabel:
    def test_skipped_job_is_not_done(self, monkeypatch):
        from bot.services import jobs as jobs_mod
        from bot.services.jobs import JobManager

        finished = []

        class _Reporter:
            async def start(self):
                return None

            async def finish(self, text):
                finished.append(text)

        async def pool_run(*a, only=None):
            return False  # RISKY job, no workers

        async def scenario():
            manager = JobManager()
            await manager.run(None, 1, ns(workers=[], run=pool_run), ns(), None,
                              lambda f, r: None, "job", "Готово ✅")
            await asyncio.sleep(0.01)

        monkeypatch.setattr(jobs_mod, "TelegramReporter", lambda *a, **k: _Reporter())
        asyncio.run(scenario())
        assert finished == ["⚠️ Не выполнено"]

    def test_skipped_job_runs_its_cleanup(self, monkeypatch):
        # the job's own finally (content.cleanup) never runs when pool.run skips it:
        # the broadcast's temp media must still go
        from bot.services import jobs as jobs_mod
        from bot.services.jobs import JobManager

        cleaned = []

        class _Reporter:
            async def start(self):
                return None

            async def finish(self, text):
                pass

        async def pool_run(*a, only=None):
            return False  # RISKY job, the only worker is scraping

        async def scenario():
            manager = JobManager()
            await manager.run(None, 1, ns(workers=[], run=pool_run), ns(), None,
                              lambda f, r: None, "job", "Готово ✅",
                              cleanup=lambda: cleaned.append(True))
            await asyncio.sleep(0.01)

        monkeypatch.setattr(jobs_mod, "TelegramReporter", lambda *a, **k: _Reporter())
        asyncio.run(scenario())
        assert cleaned == [True]


class TestReadTableKeepsNaText:
    def test_xlsx_and_csv(self, tmp_path):
        from scraper.datafiles import read_table, save_table

        df = pd.DataFrame({"Name": ["NA", "Nan", None], "n": [1, 2, 3]})
        for fmt in ("excel", "csv"):
            out = read_table(save_table(df, tmp_path / f"t_{fmt}", fmt))
            assert list(out["Name"][:2]) == ["NA", "Nan"] and pd.isna(out["Name"][2])


# --- round 6 -----------------------------------------------------------------

class TestInvitingPublicSource:
    def test_join_result_unwrapped(self):
        from functions.inviting import InvitingFunc

        chat = ns(id=4)

        class _S:
            async def __call__(self, request):
                return ns(updates=ns(chats=[chat]))  # JoinChannel -> ChatInviteJoinResultOk

        fn = InvitingFunc(_Storage(), ns(delay=[0]))
        assert asyncio.run(fn.resolve_source(_S(), "@public_chat")) is chat

    def test_join_result_without_chats_falls_back_to_a_lookup(self):
        from telethon.tl.functions.messages import CheckChatInviteRequest
        from telethon.tl.types import UpdatesTooLong
        from functions.inviting import InvitingFunc

        public, private = ns(id=4), ns(id=5)

        class _S:  # joined, but the result is UpdatesTooLong: no chats in it (as the joiner handles)
            async def __call__(self, request):
                if isinstance(request, CheckChatInviteRequest):
                    return ns(chat=private)  # ChatInviteAlready
                return ns(updates=UpdatesTooLong())

            async def get_entity(self, ref):
                assert ref == "@public_chat"
                return public

        fn = InvitingFunc(_Storage(), ns(delay=[0]))
        assert asyncio.run(fn.resolve_source(_S(), "@public_chat")) is public
        assert asyncio.run(fn.resolve_source(_S(), "https://t.me/+abcdef")) is private

    def test_basic_group_destination_stops_once(self):
        from functions.inviting import InvitingFunc

        users = [ns(id=i, bot=False, deleted=False, is_self=False) for i in (1, 2, 3)]

        from telethon import types as tl

        calls = []

        class _S:
            async def get_me(self):
                return ns(first_name="A")

            async def get_entity(self, dest):
                return tl.Chat(id=1, title="basic", photo=None, participants_count=3, date=None, version=1)

            async def iter_participants(self, source):
                for u in users:
                    yield u

            async def __call__(self, request):
                calls.append(request)

        fn = InvitingFunc(_Storage(), ns(delay=[0]))
        fn.delay = _no_sleep

        async def resolve_source(session, link):
            return "src"

        fn.resolve_source = resolve_source
        msgs, report = collect()
        asyncio.run(fn.invite(_S(), "src", "dest", [1, 2, 3], report))
        assert msgs == ["[!] [A] чат назначения — не супергруппа и не канал"] and calls == []


class TestBasicGroupAdmins:
    def test_admins_by_participant_type(self):
        from telethon import types as tl
        from functions.broadcast import Broadcast

        members = [ns(id=1, participant=tl.ChatParticipant(user_id=1, inviter_id=0, date=None)),
                   ns(id=2, participant=tl.ChatParticipantAdmin(user_id=2, inviter_id=0, date=None)),
                   ns(id=3, participant=tl.ChatParticipantCreator(user_id=3))]
        mentioned = []

        class _S:
            async def get_me(self):
                return ns(first_name="A")

            async def get_participants(self, peer, filter=None):
                return members  # a basic group: Telethon ignores the filter

        fn = Broadcast(_Storage(), ns(delay=[0], messages_count=1, messages=["hi"]))
        fn.configure(choice=0, mention_all=True, mention_mode="admins", content=ns(text="hi", entities=[], media=[], sent_media={}))
        fn.delay = _no_sleep

        async def send(session, peer, content, report, reply_to=None):
            mentioned.extend(e.user_id for e in content.entities)

        fn._send = send
        _, report = collect()
        asyncio.run(fn.broadcast(_S(), "chat", report))
        assert sorted(mentioned) == [2, 3]


class TestHeartReaction:
    def test_variation_selector_stripped(self):
        from functions.reactions import ReactionsFunc

        sent = []

        class _S:
            async def get_me(self):
                return ns(first_name="A")

            async def __call__(self, request):
                sent.append(request.reaction[0].emoticon)

        fn = ReactionsFunc(_Storage(), ns(delay=[0]))
        _, report = collect()
        asyncio.run(fn.set_reaction(_S(), "chan", 1, report, reaction="❤️"))
        assert sent == ["❤"]


class TestCaptchaHandler:
    def _run(self, buttons):
        from functions.joiner import JoinerFunc

        clicked = []

        async def click(data=None):
            clicked.append(data)

        async def get_buttons():
            return buttons

        msg = ns(mentioned=True, get_buttons=get_buttons, click=click)
        asyncio.run(JoinerFunc(_Storage(), ns(delay=[0])).on_message(msg))
        return clicked

    def test_url_or_reply_button_ignored(self):
        assert self._run([[ns(data=None)]]) == []  # MessageButton.data is None there
        assert self._run(None) == []

    def test_callback_button_clicked_with_raw_bytes(self):
        data = b"\xff\x01captcha"  # not UTF-8
        assert self._run([[ns(data=data)]]) == [data]

    def test_real_layer_button_data(self):
        from telethon.tl import types as tl
        from telethon.tl.custom.messagebutton import MessageButton

        # the installed layer: KeyboardButton(text, type=InlineButtonTypeCallback(data))
        button = tl.KeyboardButton("go", tl.InlineButtonTypeCallback(data=b"\x01x"))
        assert MessageButton(None, button, None, None, 1).data == b"\x01x"

    def test_returns_whether_clicked(self):
        from functions.joiner import JoinerFunc

        async def click(data=None):
            pass

        def msg(mentioned, buttons):
            async def get_buttons():
                return buttons
            return ns(mentioned=mentioned, get_buttons=get_buttons, click=click)

        fn = JoinerFunc(_Storage(), ns(delay=[0]))
        assert asyncio.run(fn.on_message(msg(True, [[ns(data=b"x")]]))) is True
        assert asyncio.run(fn.on_message(msg(True, [[ns(data=None)]]))) is False
        assert asyncio.run(fn.on_message(msg(False, [[ns(data=b"x")]]))) is False


class TestJoinerBotCaptcha:
    class _S:
        """Worker session: joins fine; optionally posts a captcha right after the join."""

        def __init__(self, captcha):
            self.captcha, self.handlers, self.clicked = captcha, [], []

        def add_event_handler(self, cb, ev):
            self.handlers.append(cb)

        def remove_event_handler(self, cb, ev):
            self.handlers.remove(cb)

        async def __call__(self, request):
            if self.captcha:
                async def get_buttons():
                    return [[ns(data=b"ok")]]

                async def click(data=None):
                    self.clicked.append(data)

                msg = ns(mentioned=True, get_buttons=get_buttons, click=click)
                for handler in list(self.handlers):
                    asyncio.ensure_future(handler(msg))
            return ns(chats=[ns(id=5)])

    def _run(self, monkeypatch, session, captcha):
        from functions import joiner

        monkeypatch.setattr(joiner, "CAPTCHA_WAIT", 0.05)
        fn = joiner.JoinerFunc(_Storage(), ns(delay=[0]))
        fn.sessions = [session]
        msgs, report = collect()
        asyncio.run(fn.run("1", "@chat", [0], report, captcha=captcha))
        return msgs

    def test_captcha_clicked(self, monkeypatch):
        s = self._S(captcha=True)
        msgs = self._run(monkeypatch, s, captcha=True)
        assert s.clicked == [b"ok"]
        assert any("капча пройдена" in m for m in msgs)
        assert "Итого: 1/1 аккаунтов" in msgs and "[аккаунт 1] вступил" in msgs
        assert s.handlers == []  # removed before the session is released

    def test_no_captcha_times_out(self, monkeypatch):
        s = self._S(captcha=False)
        msgs = self._run(monkeypatch, s, captcha=True)
        assert any("капчи не было" in m for m in msgs)
        assert "Итого: 1/1 аккаунтов" in msgs and "[аккаунт 1] вступил" in msgs
        assert s.handlers == []

    def test_disabled_registers_no_handler(self, monkeypatch):
        s = self._S(captcha=True)
        msgs = self._run(monkeypatch, s, captcha=False)
        assert s.clicked == []
        assert not any("captcha" in m for m in msgs)


class TestUpdaterPipFailure:
    def test_reported_without_restart(self, monkeypatch):
        import subprocess
        from modules import updater

        new_head = object()
        old_head = ns(diff=lambda other: [ns(b_path="requirements.txt")])
        repo = ns(head=ns(commit=old_head))
        repo.remote = lambda name: ns(pull=lambda: setattr(repo.head, "commit", new_head))
        monkeypatch.setattr(updater, "Repo", lambda *_: repo)

        def fail(console):
            raise subprocess.CalledProcessError(1, ["pip"])

        monkeypatch.setattr(updater, "update_requirements", fail)
        restarted, printed = [], []
        monkeypatch.setattr(updater, "restart_app", lambda: restarted.append(1))
        console = ns(status=lambda *_: contextlib.nullcontext(), print=lambda *a: printed.append(a[0]))

        updater.update(console)
        assert restarted == [] and any("Не удалось установить зависимости" in p for p in printed)


class TestReporterFloodRetries:
    def test_final_survives_two_retry_afters(self, monkeypatch):
        from aiogram.exceptions import TelegramRetryAfter
        from bot.services import runner as runner_mod
        from bot.services.runner import TelegramReporter

        monkeypatch.setattr(runner_mod.asyncio, "sleep", _no_sleep)
        failures = [TelegramRetryAfter(method=None, message="flood", retry_after=1)] * 2
        sent = []

        class _Bot:
            async def edit_message_text(self, text, **k):
                if failures:
                    raise failures.pop()
                sent.append(text)

        r = TelegramReporter(_Bot(), 1, header="H")
        r.message_id = 1
        asyncio.run(r.finish("Готово ✅"))
        assert sent and sent[-1].startswith("Готово ✅")

    def test_progress_retry_after_schedules_flush(self):
        from aiogram.exceptions import TelegramRetryAfter
        from bot.services.runner import TelegramReporter

        failures = [TelegramRetryAfter(method=None, message="flood", retry_after=0)]
        sent = []

        class _Bot:
            async def edit_message_text(self, text, **k):
                if failures:
                    raise failures.pop()
                sent.append(text)

        async def scenario():
            r = TelegramReporter(_Bot(), 1, header="H", min_interval=0.01)
            r.message_id = 1
            await r("tail line")  # hits flood control, held back
            await asyncio.sleep(0.1)

        asyncio.run(scenario())
        assert sent and "tail line" in sent[-1]


class TestReportAnswerFailure:
    def test_busy_not_stuck(self):
        from bot.routers import moderation

        stepped = []

        async def step(flow, option):
            stepped.append(option)
            return "choose", [ns(option=b"b", text="b")]

        async def failing_answer(*a, **k):
            raise RuntimeError("query is too old")

        async def answer(*a, **k):
            return None

        flow = {"busy": False, "selections": []}
        moderation._FLOWS[7] = (ns(step=step), flow, [ns(option=b"a")])
        cb = ns(message=ns(chat=ns(id=7), answer=answer), answer=failing_answer, bot=None)
        try:
            asyncio.run(moderation.rm_choose(cb, ns(value="0"), _Manager(), ns(in_job=[])))
            assert stepped == [b"a"] and moderation._FLOWS[7][1]["busy"] is False
        finally:
            moderation._FLOWS.clear()


class TestSnapshotDanglingSymlink:
    def test_skipped(self, tmp_path):
        import os
        from bot.services.scraping import dir_snapshot

        (tmp_path / "real.parquet").write_text("x")
        os.symlink(tmp_path / "gone.parquet", tmp_path / "latest.parquet")
        assert list(dir_snapshot(str(tmp_path))) == ["real.parquet"]


class TestReadPreviewUtf16:
    def test_emoji_preview_fits(self):
        from bot.routers.scraping import read_preview

        out = read_preview(pd.DataFrame({"c": ["😀" * 1000] * 10}))
        body = out[len("<pre>"):out.index("</pre>")]
        assert len(html.unescape(body).encode("utf-16-le")) // 2 <= 3501


class TestCaptureTimeout:
    def test_download_gets_long_timeout(self, tmp_path):
        from bot.services.capture import _download

        seen = {}

        class _Bot:
            async def download(self, source, destination=None, timeout=30):
                seen["timeout"] = timeout

        asyncio.run(_download(_Bot(), str(tmp_path), 0, ("src", "a.mp4", False, False, False, False)))
        assert seen["timeout"] >= 300


class TestStopWhileFinishing:
    def test_completed_job_not_relabelled(self, monkeypatch):
        from bot.services import jobs as jobs_mod
        from bot.services.jobs import JobManager

        finished = []
        manager = JobManager()

        class _Reporter:
            async def start(self):
                return None

            async def finish(self, text):
                # Stop pressed while the final status waits out flood control
                stopped.append(await manager.stop())
                await asyncio.sleep(0.01)
                finished.append(text)

        stopped = []

        async def pool_run(inst, bf, fac, rep, only=None):
            return True

        async def scenario():
            await manager.run(None, 1, ns(workers=[], run=pool_run), ns(), None,
                              lambda f, r: None, "job", "Готово ✅")
            await asyncio.sleep(0.05)

        monkeypatch.setattr(jobs_mod, "TelegramReporter", lambda *a, **k: _Reporter())
        asyncio.run(scenario())
        assert stopped == [False] and finished == ["Готово ✅"] and manager.active is False



class TestReportFlowHoldsItsWorkers:
    def test_busy_until_aborted(self):
        from bot.routers import moderation

        class _First:
            async def connect(self):
                return None

            async def disconnect(self):
                return None

        async def step(flow, option):
            return "choose", [ns(text="spam", option=b"1")]

        class _Manager:
            label = ""

            def acquire(self, *a, **k):
                return True

            def release(self):
                return None

        class _State:
            async def get_data(self):
                return {"link": "l", "ids": [1]}

            async def clear(self):
                return None

        async def answer(*a, **k):
            return None

        workers = [_First(), _First()]
        pool = ns(delegate=lambda inst: workers, in_job=[])
        msg = ns(text="-", chat=ns(id=7), answer=answer, bot=None)
        try:
            asyncio.run(moderation.rm_begin(msg, _State(), pool, {"ReportFunc": ns(step=step)}, _Manager()))
            assert pool.in_job == workers  # the flow drives them: no scrape may start on them
            asyncio.run(moderation._abort(7, pool))  # /cancel or the inactivity timeout
            assert pool.in_job == []
        finally:
            moderation._FLOWS.clear()


# --- a dead worker: Telethon's get_me() returns None for a banned / logged-out account ---

class _Worker:
    def __init__(self, me):
        self.me, self.requests = me, []

    async def get_me(self):
        return self.me

    async def __call__(self, request):
        self.requests.append(request)

    async def send_message(self, *args, **kwargs):
        self.requests.append(args)

    async def get_input_entity(self, peer):  # a private link's chat, from the cache
        return peer


class TestDeadWorkerIsSkipped:
    def _run(self, cls, call, settings):
        dead, alive = _Worker(None), _Worker(ns(id=1, first_name="alive"))
        fn = cls(_Storage([dead, alive]), settings)
        msgs, report = collect()
        asyncio.run(call(fn, report))
        return dead, alive, msgs

    def test_profile_change(self):
        from functions.changebio import ChangeBioFunc

        dead, alive, msgs = self._run(ChangeBioFunc, lambda f, r: f.run(r, bio="x"), ns(profile_pause=[0]))
        assert not dead.requests and len(alive.requests) == 1
        assert any("не удалось опросить аккаунт" in m for m in msgs)

    def test_reactions(self):
        from functions.reactions import ReactionsFunc

        dead, alive, msgs = self._run(ReactionsFunc, lambda f, r: f.run("https://t.me/c/1/2", "👍", r),
                                      ns(delay=[0]))
        assert not dead.requests and len(alive.requests) == 1

    def test_comments(self):
        from functions.broadcast_comments import CommentsBroadcastFunc
        from modules.rich_message import RichContent

        dead, alive, msgs = self._run(
            CommentsBroadcastFunc,
            lambda f, r: f.run("https://t.me/name/5", RichContent(text="hi"), [0], r),
            ns(messages_count=1),
        )
        assert not dead.requests and len(alive.requests) == 1


# --- inviting: link forms -----------------------------------------------------------

@pytest.mark.parametrize("link, public, ref", [
    ("https://t.me/name/", True, "@name"),
    ("https://t.me/name?start=1", True, "@name"),
    ("t.me/name/123", True, "@name"),
    ("@name", True, "@name"),
    ("https://t.me/+AbC/", False, "AbC"),
    ("https://t.me/joinchat/AbC", False, "AbC"),
    ("+AbC", False, "AbC"),
])
def test_inviting_link_forms(link, public, ref):
    from functions.inviting import InvitingFunc

    assert InvitingFunc.is_public(link) is public
    assert (InvitingFunc.public_ref(link) if public else InvitingFunc.invite_hash(link)) == ref


# --- stickers campaign: the set is fetched once per worker ------------------------------

def test_sticker_set_fetched_once_per_worker():
    from telethon.tl.functions.messages import GetStickerSetRequest
    from functions.broadcast import Broadcast
    from modules.rich_message import RichContent

    class _Chat(_Worker):
        async def __call__(self, request):
            self.requests.append(request)
            return ns(documents=["sticker"])

        async def send_file(self, *args, **kwargs):
            self.requests.append(args)

    session = _Chat(ns(id=1, first_name="w"))
    fn = Broadcast(_Storage([session]), ns(messages_count=3, delay=[0]))
    fn.configure(4, sticker_set="https://t.me/addstickers/Pack", delay=[0], content=RichContent(text="hi"))
    msgs, report = collect()

    asyncio.run(fn.broadcast(session, "chat", report))
    assert sum(isinstance(r, GetStickerSetRequest) for r in session.requests) == 1
    assert sum("отправлено, всего" in m for m in msgs) == 3


# --- joiner: link forms -------------------------------------------------------------

@pytest.mark.parametrize("link, request_name, arg", [
    ("https://t.me/name/", "JoinChannelRequest", "@name"),
    ("t.me/name/123", "JoinChannelRequest", "@name"),
    ("https://t.me/name?x=1", "JoinChannelRequest", "@name"),
    ("@name", "JoinChannelRequest", "@name"),
    ("https://t.me/joinchat/AbC/", "ImportChatInviteRequest", "AbC"),
    ("https://t.me/+AbC", "ImportChatInviteRequest", "AbC"),
])
def test_joiner_link_forms(link, request_name, arg):
    from functions.joiner import JoinerFunc

    session = _Worker(ns(id=1, first_name="w"))
    fn = JoinerFunc(_Storage([session]), ns(delay=[0]))
    msgs, report = collect()

    asyncio.run(fn.join(session, link, 0, "1", report))
    request = session.requests[0]
    assert type(request).__name__ == request_name
    assert (request.channel if request_name == "JoinChannelRequest" else request.hash) == arg


# --- clear dialogs: a worker whose dialog list fails doesn't end the job --------------

def test_clear_dialogs_survives_a_failed_listing():
    from functions.clear_chats import ClearDialogsFunc

    class _Dialogs(_Worker):
        def __init__(self, me, fail):
            super().__init__(me)
            self.fail = fail

        async def iter_dialogs(self):
            if self.fail:
                raise ConnectionError("dropped")
            yield ns(entity=ns(), id=5, title="chat")

        async def __call__(self, request):
            self.requests.append(request)
            return ns(offset=0)

    broken, ok = _Dialogs(ns(id=1, first_name="a"), True), _Dialogs(ns(id=2, first_name="b"), False)
    fn = ClearDialogsFunc(_Storage([broken, ok]), ns())
    msgs, report = collect()

    asyncio.run(fn.run(report))
    assert any("не удалось получить список диалогов" in m for m in msgs)
    assert len(ok.requests) == 1 and any("удалён" in m for m in msgs)


# --- a private post link (t.me/c/<id>/<msg>): its chat resolves from the cache, else the dialogs ---

class _Resolver:
    def __init__(self, cached):
        self.cached, self.dialogs_loaded = cached, 0

    async def get_input_entity(self, peer):
        if not self.cached:
            raise ValueError("Could not find the input entity")
        return ("input", peer.channel_id)

    async def get_dialogs(self):
        self.dialogs_loaded += 1
        self.cached = True


def _resolve(session, peer):
    from functions.base import TelethonFunction

    return asyncio.run(TelethonFunction(_Storage(), ns()).resolve_chat(session, peer))


def test_a_private_link_chat_loads_the_dialogs_once_on_a_cache_miss():
    from telethon.tl.types import PeerChannel

    session = _Resolver(cached=False)
    assert _resolve(session, PeerChannel(5)) == ("input", 5) and session.dialogs_loaded == 1


def test_a_cached_private_link_chat_needs_no_dialogs():
    from telethon.tl.types import PeerChannel

    session = _Resolver(cached=True)
    assert _resolve(session, PeerChannel(5)) == ("input", 5) and session.dialogs_loaded == 0


def test_a_public_peer_is_passed_as_is():
    session = _Resolver(cached=False)
    assert _resolve(session, "durov") == "durov" and session.dialogs_loaded == 0


# --- the «Очистить диалоги» confirmation expires: an old «Да» must not wipe the workers ---

def _clear(age):
    import datetime

    from bot.callbacks import ChoiceCB
    from bot.routers import service

    alerts, runs = [], []

    async def answer(text=None, show_alert=False):
        if show_alert:
            alerts.append(text)

    class _Manager:
        async def run(self, *args, **kwargs):
            runs.append(args)

    message = ns(date=datetime.datetime.now(datetime.timezone.utc) - age, chat=ns(id=1))
    callback = ns(message=message, answer=answer, bot=None)
    functions = {"ClearDialogsFunc": object()}
    asyncio.run(service.clear_run(callback, ChoiceCB(scope="clear_confirm", value="yes"),
                                  ns(), functions, _Manager()))
    return alerts, runs


def test_a_stale_clear_confirmation_runs_nothing():
    import datetime

    alerts, runs = _clear(datetime.timedelta(hours=1))
    assert not runs and alerts and "устарело" in alerts[0]


def test_a_fresh_clear_confirmation_runs_the_wipe():
    import datetime

    alerts, runs = _clear(datetime.timedelta(seconds=5))
    assert runs and not alerts
